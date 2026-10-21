"""Generate privacy-safe SUIM screenshots and a matching demo workbook."""
import argparse
from pathlib import Path
from typing import Iterable, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image, ImageDraw, ImageFont


WIDTH, HEIGHT = 1600, 1100
BG = "#F3F5F7"
NAVY = "#18324B"
BLUE = "#47799A"
INK = "#18232D"
MUTED = "#596A78"
LINE = "#B9C5CE"
WHITE = "#FFFFFF"
FIELD_DESCRIPTIONS = {
    "ACTVT": "Activity",
    "TCD": "Transaction Code",
}


def _font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    """Load a Windows UI font, falling back to Pillow's built-in font."""
    candidates = (
        Path("C:/Windows/Fonts/segoeuib.ttf" if bold else "C:/Windows/Fonts/segoeui.ttf"),
        Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
    )
    for candidate in candidates:
        if candidate.is_file():
            return ImageFont.truetype(str(candidate), size)
    return ImageFont.load_default()


def _draw_screen(path: Path, objects: Iterable[Tuple[str, str, str]]) -> None:
    """Render a synthetic SUIM authorization-criteria screen."""
    image = Image.new("RGB", (WIDTH, HEIGHT), BG)
    draw = ImageDraw.Draw(image)
    regular = _font(26)
    small = _font(21)
    bold = _font(27, bold=True)
    title = _font(34, bold=True)

    draw.rectangle((0, 0, WIDTH, 74), fill=NAVY)
    draw.text((42, 18), "SAP  |  User Information System", fill=WHITE, font=bold)
    draw.rectangle((0, 74, WIDTH, 126), fill="#E4EAF0")
    draw.text((42, 87), "Transaction: SUIM", fill=INK, font=small)
    draw.text((WIDTH - 460, 87), "SYNTHETIC TRAINING DATA", fill=BLUE, font=small)

    draw.text((50, 165), "Users by Complex Selection Criteria", fill=INK, font=title)
    draw.text((50, 220), "Authorization selection", fill=MUTED, font=regular)

    y = 275
    for obj_name, field_name, rows in objects:
        box_bottom = y + 235
        draw.rounded_rectangle((44, y, WIDTH - 44, box_bottom), radius=8,
                               fill=WHITE, outline=LINE, width=2)
        draw.rectangle((46, y + 2, WIDTH - 46, y + 58), fill="#EAF0F5")
        draw.text((68, y + 13), f"Authorization Object {obj_name}",
                  fill=NAVY, font=bold)
        description = FIELD_DESCRIPTIONS.get(field_name, field_name)
        draw.text((82, y + 82), f"{field_name} - {description}",
                  fill=INK, font=regular)
        draw.line((82, y + 122, WIDTH - 82, y + 122), fill=LINE, width=1)
        for index, value_row in enumerate(rows.split("|")):
            draw.text((112, y + 140 + index * 42), value_row.strip(),
                      fill=INK, font=regular)
        y = box_bottom + 24

    draw.rounded_rectangle((50, 920, 205, 978), radius=5, fill=BLUE)
    draw.text((89, 933), "Execute", fill=WHITE, font=bold)
    draw.text((230, 937), "All names and authorization values on this screen are fictional.",
              fill=MUTED, font=small)
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path, format="PNG")


def _write_workbook(path: Path) -> None:
    """Create one expected pass and two intentional failures."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Synthetic Demo"
    headers = (
        "RAFIT", "Function", "Auth Object 1", "Auth Object 1 Fields",
        "Auth Object 2", "Auth Object 2 Fields",
    )
    sheet.append(headers)
    rows = (
        (
            "DEMO-001", "Expected pass - all values and objects are shown",
            "S_TCODE", "TCD = SU01 OR SU10", "S_USER_GRP", "ACTVT = 01 AND 02",
        ),
        (
            "DEMO-002", "Expected fail - required transaction is absent",
            "S_TCODE", "TCD = SU99", None, None,
        ),
        (
            "DEMO-003", "Expected fail - authorization object is absent",
            "S_USER_AGR", "ACTVT = 02", None, None,
        ),
    )
    for row in rows:
        sheet.append(row)

    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3864")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column, width in {
        "A": 16, "B": 58, "C": 20, "D": 30, "E": 20, "F": 30,
    }.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def generate(output_dir: Path) -> None:
    """Write a demo workbook and its synthetic SAP-style screenshot."""
    output_dir.mkdir(parents=True, exist_ok=True)
    _write_workbook(output_dir / "demo_master.xlsx")
    _draw_screen(
        output_dir / "demo_suim.png",
        (
            ("S_TCODE", "TCD", "Value SU01 OR SU10"),
            ("S_USER_GRP", "ACTVT", "Value 01|AND 02"),
        ),
    )


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=Path(__file__).resolve().parents[1] / "sample_data" / "synthetic_demo",
        help="Destination directory (default: sample_data/synthetic_demo)",
    )
    args = parser.parse_args()
    generate(args.output_dir)
    print(f"Synthetic demo files created in: {args.output_dir.resolve()}")


if __name__ == "__main__":
    main()
