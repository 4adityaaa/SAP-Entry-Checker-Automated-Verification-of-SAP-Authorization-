"""Reads the master Excel (critical codes) into structured requirements.

Each row has pairs of columns:  Auth Object N | Auth Object N Fields
Fields look like:  "TCD = SU01 OR SU10"
                   "ACTVT = 02; DICBERCLS = SS"
                   "ACTVT = (01 OR 02) AND 07"
                   "TABLE = '*'"
"""
import re
from dataclasses import dataclass, field
from typing import List, Optional

from openpyxl import load_workbook


# ----------------------------------------------------------------------------
# Tiny boolean-expression parser for   (01 OR 02) AND 07
# ----------------------------------------------------------------------------
_TOKEN_RE = re.compile(r"""'[^']*'|"[^"]*"|\(|\)|[^\s()]+""")


def _tokenise(expr: str) -> List[str]:
    return _TOKEN_RE.findall(expr)


def parse_expression(expr: str):
    """Returns a nested tuple tree: ('or', [...]) / ('and', [...]) / ('val', 'X')"""
    tokens = _tokenise(expr)
    pos = 0

    def peek() -> Optional[str]:
        return tokens[pos] if pos < len(tokens) else None

    def take():
        nonlocal pos
        tok = tokens[pos]
        pos += 1
        return tok

    def parse_or():
        items = [parse_and()]
        while peek() is not None and peek().upper() == "OR":
            take()
            items.append(parse_and())
        return items[0] if len(items) == 1 else ("or", items)

    def parse_and():
        items = [parse_atom()]
        while peek() is not None and peek().upper() == "AND":
            take()
            items.append(parse_atom())
        return items[0] if len(items) == 1 else ("and", items)

    def parse_atom():
        tok = peek()
        if tok is None:
            raise ValueError("Expected a value at the end of the expression.")
        if tok == "(":
            take()
            if peek() == ")":
                raise ValueError("Empty parentheses are not valid in a field expression.")
            node = parse_or()
            if peek() != ")":
                raise ValueError("Missing closing ')' in field expression.")
            take()
            return node
        if tok == ")" or tok.upper() in {"AND", "OR"}:
            raise ValueError(f"Expected a value, found {tok!r}.")
        take()
        if len(tok) >= 2 and tok[0] == tok[-1] and tok[0] in {"'", '"'}:
            tok = tok[1:-1]
        if not tok:
            raise ValueError("Field values cannot be empty.")
        return ("val", tok)

    if not tokens:
        raise ValueError("Field expression cannot be empty.")
    tree = parse_or()
    if pos < len(tokens):
        raise ValueError(f"Unexpected token {tokens[pos]!r} in field expression.")
    return tree


def tree_values(tree) -> List[str]:
    """All distinct literal values in the tree (order kept)."""
    out: List[str] = []

    def walk(n):
        if n[0] == "val":
            if n[1] and n[1] not in out:
                out.append(n[1])
        else:
            for c in n[1]:
                walk(c)

    walk(tree)
    return out


def tree_eval(tree, present) -> bool:
    """present(value)->bool. Evaluates the boolean expression."""
    kind = tree[0]
    if kind == "val":
        return present(tree[1])
    if kind == "and":
        return all(tree_eval(c, present) for c in tree[1])
    return any(tree_eval(c, present) for c in tree[1])


# ----------------------------------------------------------------------------
@dataclass
class FieldReq:
    name: str                  # e.g. ACTVT
    raw: str                   # e.g. "(01 OR 02) AND 07"
    tree: tuple
    values: List[str]


@dataclass
class ObjectReq:
    obj: str                   # e.g. S_TCODE
    fields: List[FieldReq] = field(default_factory=list)
    raw: str = ""


@dataclass
class RowReq:
    excel_row: int
    rafit: str
    function: str
    objects: List[ObjectReq] = field(default_factory=list)


def parse_fields(text: str) -> List[FieldReq]:
    out: List[FieldReq] = []
    if text is None:
        return out
    for part in str(text).split(";"):
        part = part.strip()
        if not part:
            continue
        if "=" not in part:
            raise ValueError(f"Expected 'FIELD = value' in {part!r}.")
        name, expr = part.split("=", 1)
        name = name.strip().upper()
        expr = expr.strip()
        if not re.fullmatch(r"[A-Z][A-Z0-9_]*", name):
            raise ValueError(f"Invalid authorization field name {name!r}.")
        try:
            tree = parse_expression(expr)
        except ValueError as exc:
            raise ValueError(f"Invalid expression for field {name}: {exc}") from exc
        out.append(FieldReq(name=name, raw=expr, tree=tree, values=tree_values(tree)))
    return out


def _clean(v) -> str:
    return "" if v is None else str(v).strip()


def load_requirements(path: str) -> List[RowReq]:
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        raise ValueError("The Excel file is empty.")

    header = [_clean(h) for h in rows[0]]
    # locate (object, fields) column pairs by header text; fall back to position
    pairs = []
    for i, h in enumerate(header):
        hl = h.lower()
        if hl.startswith("auth object") and "field" not in hl and i + 1 < len(header):
            pairs.append((i, i + 1))
    if not pairs:
        pairs = [(i, i + 1) for i in range(2, len(header) - 1, 2)]

    rafit_col = next((i for i, h in enumerate(header) if h.lower().startswith("rafit")), 0)
    func_col = next((i for i, h in enumerate(header) if h.lower().startswith("function")), 1)

    reqs: List[RowReq] = []
    for r_idx, row in enumerate(rows[1:], start=2):
        row = list(row) + [None] * (len(header) - len(row))
        if all(_clean(c) == "" for c in row):
            continue
        rr = RowReq(excel_row=r_idx, rafit=_clean(row[rafit_col]), function=_clean(row[func_col]))
        for oi, fi in pairs:
            obj = _clean(row[oi]).upper()
            if not obj:
                continue
            raw = _clean(row[fi])
            rr.objects.append(ObjectReq(obj=obj, fields=parse_fields(raw), raw=raw))
        if rr.objects:
            reqs.append(rr)
    if not reqs:
        raise ValueError("No authorization-object rows were found in the Excel file. "
                         "Expected columns like 'Auth Object 1' and 'Auth Object 1 Fields'.")
    return reqs


def vocabulary(reqs: List[RowReq]):
    """Known objects / fields / values - used to guide the OCR parsing."""
    objs, flds, vals = set(), set(), set()
    for r in reqs:
        for o in r.objects:
            objs.add(o.obj)
            for f in o.fields:
                flds.add(f.name)
                vals.update(f.values)
    return objs, flds, vals
