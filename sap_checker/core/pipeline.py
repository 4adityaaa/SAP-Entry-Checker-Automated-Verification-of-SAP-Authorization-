"""End-to-end processing: screenshots + Excel -> results (+ Excel report)."""
import os
import re
import time
from typing import Callable, List, Tuple

import cv2
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .excel_parser import load_requirements, vocabulary
from .extractor import extract_blocks
from .matcher import match_all
from .ocr_engine import read_screenshot


def natural_key(s: str):
    return [(0, int(t)) if t.isdigit() else (1, t.lower())
            for t in re.split(r"(\d+)", s)]


def _annotate(img, extracted, out_path, max_w=1100):
    vis = img.copy()
    used = set(extracted["used"])
    for it in extracted["items"]:
        pass
    # re-draw boxes: we only kept centres, so draw rectangles from stored geometry
    for it in extracted["items"]:
        x0, cy, h = int(it["x0"]), int(it["cy"]), int(it["h"])
        w = int(max(2 * (it["cx"] - it["x0"]), 4))
        color = (0, 170, 0) if it["idx"] in used else (150, 150, 150)
        thick = 3 if it["idx"] in used else 1
        cv2.rectangle(vis, (x0, cy - h // 2), (x0 + w, cy + h // 2), color, thick)
    scale = min(1.0, max_w / vis.shape[1])
    if scale < 1.0:
        vis = cv2.resize(vis, (int(vis.shape[1] * scale), int(vis.shape[0] * scale)),
                         interpolation=cv2.INTER_AREA)
    cv2.imwrite(out_path, vis, [cv2.IMWRITE_JPEG_QUALITY, 80])


def _merge_continuations(screens: List[dict]):
    """If a screenshot starts mid-block (page scrolled), attach those fields to
    the last block of the previous screenshot."""
    previous_block = None
    for screen in screens:
        blocks = screen["blocks"]
        if blocks and blocks[0]["object"] is None and previous_block is not None:
            continuation = blocks.pop(0)
            for field_name, values in continuation["fields"].items():
                target_values = previous_block["fields"].setdefault(field_name, [])
                for value in values:
                    if value not in target_values:
                        target_values.append(value)
            evidence = previous_block.setdefault(
                "evidence_screens", [previous_block["screenshot"]]
            )
            if screen["filename"] not in evidence:
                evidence.append(screen["filename"])
            screen["notes"].append(
                f"Continued block {previous_block['object']} from previous screenshot"
            )

        object_blocks = [block for block in blocks if block["object"]]
        for block in object_blocks:
            block.setdefault("screenshot", screen["filename"])
            block.setdefault("evidence_screens", [screen["filename"]])
        if object_blocks:
            previous_block = object_blocks[-1]


def run_job(files: List[Tuple[str, bytes]], excel_path: str, out_dir: str,
            strict: bool, progress: Callable[[int, int, str], None]) -> dict:
    os.makedirs(out_dir, exist_ok=True)
    reqs = load_requirements(excel_path)
    objs, flds, _vals = vocabulary(reqs)
    values_by_field = {}
    for row in reqs:
        for obj in row.objects:
            for field_req in obj.fields:
                values_by_field.setdefault(field_req.name, set()).update(field_req.values)

    files = sorted(files, key=lambda f: natural_key(f[0]))
    screens: List[dict] = []
    confidence_sum = 0.0
    confidence_count = 0
    total = len(files)
    for n, (name, data) in enumerate(files, 1):
        progress(n - 1, total, f"Reading {name}")
        t0 = time.time()
        scr = {"filename": name, "blocks": [], "lines": [], "notes": [], "warnings": [],
               "error": "", "image": "", "skew": 0.0, "blur": 0.0,
               "ocr_confidence": None}
        try:
            img, items, info = read_screenshot(data)
            ex = extract_blocks(items, objs, flds, _vals, values_by_field)
            scr.update(blocks=ex["blocks"], lines=ex["lines"], warnings=info["warnings"],
                       skew=info["skew"], blur=info["blur"],
                       ocr_confidence=(round(ex["confidence"] * 100, 1)
                                       if ex["confidence"] is not None else None))
            if ex["confidence"] is not None:
                confidence_sum += ex["confidence"] * ex["confidence_count"]
                confidence_count += ex["confidence_count"]
            img_name = f"shot_{n:02d}.jpg"
            _annotate(img, ex, os.path.join(out_dir, img_name))
            scr["image"] = img_name
            if not ex["blocks"]:
                scr["warnings"].append("No authorization-object block recognised on this screenshot.")
        except Exception as exc:                       # never let one bad file stop the job
            scr["error"] = f"{type(exc).__name__}: {exc}"
        scr["seconds"] = round(time.time() - t0, 1)
        screens.append(scr)
    progress(total, total, "Matching with Excel")

    _merge_continuations(screens)
    results = match_all(reqs, screens, strict=strict)
    passed = sum(1 for r in results if r["status"] == "PASS")
    checks = []
    for row in results:
        for object_result in row["objects"]:
            checks.extend(object_result["fields"] or [{"ok": object_result["ok"]}])
    if checks:
        criteria_total = len(checks)
        criteria_matched = sum(1 for result in checks if result["ok"])
    else:
        criteria_total = 0
        criteria_matched = 0
    summary = {"total": len(results), "passed": passed, "failed": len(results) - passed,
               "screenshots": total, "mode": "strict" if strict else "lenient",
               "criteria_matched": criteria_matched, "criteria_total": criteria_total,
               "criteria_match_percent": (
                   round(criteria_matched * 100 / criteria_total, 1)
                   if criteria_total else None
               ),
               "ocr_confidence_percent": (
                   round(confidence_sum * 100 / confidence_count, 1)
                   if confidence_count else None
               )}
    report_path = os.path.join(out_dir, "SAP_Entry_Check_Report.xlsx")
    write_report(report_path, summary, results, screens)
    return {"summary": summary, "results": results, "screens": screens}


# ----------------------------------------------------------------------------
GREEN = PatternFill("solid", fgColor="C6EFCE")
RED = PatternFill("solid", fgColor="FFC7CE")
HEAD = PatternFill("solid", fgColor="1F3864")


def _describe(o: dict) -> str:
    parts = []
    for f in o["fields"]:
        mark = "OK" if f["ok"] else "MISSING " + ",".join(f["missing"])
        parts.append(f"{f['name']} [{f['expected']}] -> {mark}")
    base = f"{o['obj']} ({'OK' if o['ok'] else 'FAIL'})"
    screenshots = o.get("screenshots") or ([o["screenshot"]] if o["screenshot"] else [])
    if screenshots:
        base += f" in {', '.join(screenshots)}"
    return base + (": " + "; ".join(parts) if parts else "")


def write_report(path: str, summary: dict, results: List[dict], screens: List[dict]):
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Total entries (Excel rows)", summary["total"]])
    ws.append(["PASS", summary["passed"]])
    ws.append(["FAIL", summary["failed"]])
    ws.append(["Screenshots processed", summary["screenshots"]])
    ws.append(["Match mode", summary["mode"]])
    ws.append(["Workbook criteria matched",
               f"{summary['criteria_matched']} / {summary['criteria_total']}"])
    ws.append(["Criteria accuracy (%)", summary["criteria_match_percent"]
               if summary["criteria_match_percent"] is not None else "N/A"])
    ws.append(["Relevant OCR confidence (%)", summary["ocr_confidence_percent"]
               if summary["ocr_confidence_percent"] is not None else "N/A"])
    ws.column_dimensions["A"].width = 30
    for r in ws.iter_rows():
        r[0].font = Font(bold=True)

    ws2 = wb.create_sheet("Results")
    ws2.append(["Excel Row", "RAFIT", "Function", "Status", "Screenshot(s)", "Details", "Reason"])
    for c in ws2[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD
    for r in results:
        ws2.append([r["excel_row"], r["rafit"], r["function"], r["status"],
                    ", ".join(r["evidence"]), "\n".join(_describe(o) for o in r["objects"]),
                    r["reason"]])
        ws2.cell(ws2.max_row, 4).fill = GREEN if r["status"] == "PASS" else RED
    for col, w in zip("ABCDEFG", (10, 14, 48, 10, 32, 90, 60)):
        ws2.column_dimensions[col].width = w
    for row in ws2.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws2.freeze_panes = "A2"

    ws3 = wb.create_sheet("Extracted from screenshots")
    ws3.append(["Screenshot", "Authorization Object", "Field", "Values read by OCR",
                "Relevant OCR confidence (%)"])
    for c in ws3[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD
    for s in screens:
        if s["error"]:
            ws3.append([s["filename"], "ERROR", s["error"], "", ""])
        for b in s["blocks"]:
            if not b["fields"]:
                ws3.append([s["filename"], b["object"], "", "", s["ocr_confidence"]])
            for f, v in b["fields"].items():
                ws3.append([s["filename"], b["object"], f, ", ".join(v),
                            s["ocr_confidence"]])
    for col, w in zip("ABCDE", (34, 24, 16, 50, 28)):
        ws3.column_dimensions[col].width = w
    wb.save(path)
